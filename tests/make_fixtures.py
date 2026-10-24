# SPDX-License-Identifier: GPL-3.0-or-later
"""Builds the PDFs the tests use. Run with: python3 tests/make_fixtures.py"""
from pathlib import Path
import subprocess

from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

out = Path(__file__).parent / "fixtures"
out.mkdir(exist_ok=True)

pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))


def sample(path):
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Inkleaf sample")
    # page 1: standard font and an embedded TrueType font
    c.setFont("Helvetica", 14)
    c.drawString(72, 760, "Invoice number: INV-2041 for Hotel Aurora")
    c.setFont("DejaVu", 12)
    c.drawString(72, 720, "Guest: Priya Sharma, Room 304, three nights")
    c.drawString(72, 700, "Total due: 18,450 rupees")
    c.setFillColorRGB(0.85, 0.92, 1)
    c.rect(60, 600, 300, 60, fill=1, stroke=0)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Times-Roman", 16)
    c.drawString(72, 625, "Text on a tinted background")
    c.showPage()
    # page 2: rotated text and a second line
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 760, "Page two heading")
    c.saveState()
    c.translate(300, 400)
    c.rotate(30)
    c.setFont("Helvetica", 12)
    c.drawString(0, 0, "Slanted line of text")
    c.restoreState()
    c.showPage()
    # page 3
    c.setFont("Courier", 12)
    c.drawString(72, 760, "Third page in Courier")
    c.showPage()
    c.save()


def form(path):
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica", 12)
    c.drawString(72, 780, "Guest registration")
    c.drawString(72, 740, "Name")
    c.acroForm.textfield(name="guest_name", x=160, y=732, width=250, height=20)
    c.drawString(72, 700, "Nights")
    c.acroForm.textfield(name="nights", x=160, y=692, width=60, height=20)
    c.drawString(72, 660, "Breakfast included")
    c.acroForm.checkbox(name="breakfast", x=220, y=656, size=16)
    c.drawString(72, 620, "Room type")
    c.acroForm.choice(name="room_type", x=160, y=612, width=150, height=20,
                      options=["Standard", "Deluxe", "Suite"], value="Standard")
    c.drawString(72, 580, "Payment")
    c.acroForm.radio(name="payment", value="card", x=160, y=576, size=14, selected=True)
    c.drawString(180, 578, "Card")
    c.acroForm.radio(name="payment", value="cash", x=240, y=576, size=14)
    c.drawString(260, 578, "Cash")
    c.showPage()
    c.save()


def photo(path):
    """Two pages sharing one big, high-quality photo, stored twice (the way
    merging a file with itself does). Compress should shrink and merge it."""
    import random
    from PIL import Image, ImageFilter
    random.seed(7)
    w, h = 2800, 1900
    img = Image.new("RGB", (w, h))
    px = img.load()
    for y in range(h):
        for x in range(w):
            px[x, y] = (int(255 * x / w), int(255 * y / h), (x * y) % 256)
    noise = Image.effect_noise((w, h), 40).convert("RGB")
    img = Image.blend(img, noise, 0.25).filter(ImageFilter.GaussianBlur(1))
    tmp = out / "_photo.jpg"
    img.save(tmp, quality=95)
    single = out / "_photo-single.pdf"
    c = canvas.Canvas(str(single), pagesize=A4)
    c.setFont("Helvetica", 18)
    c.drawString(72, 780, "Holiday photo")
    c.drawImage(str(tmp), 72, 300, width=450, height=450 * h / w)
    c.showPage()
    c.save()
    copy = out / "_photo-copy.pdf"
    copy.write_bytes(single.read_bytes())
    subprocess.run(["qpdf", "--empty", "--pages", str(single), str(copy), "--", str(path)], check=True)
    for f in (tmp, single, copy):
        f.unlink()


def scan(path):
    """A 'scanned' page: just a picture of text, no text layer, for OCR."""
    from PIL import Image, ImageDraw, ImageFont
    dpi = 200
    w, h = int(8.27 * dpi), int(11.69 * dpi)
    img = Image.new("L", (w, h), 250)
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", 44)
    lines = ["Grand Hotel Aurora", "Booking reference 58213",
             "Guest: Rohan Verma", "Check-in on Monday morning"]
    for i, line in enumerate(lines):
        d.text((180, 240 + i * 110), line, fill=20, font=font)
    tmp = out / "_scan.png"
    img.save(tmp)
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawImage(str(tmp), 0, 0, width=A4[0], height=A4[1])
    c.showPage()
    c.save()
    tmp.unlink()


def office(folder):
    """Word, Excel, PowerPoint, CSV and Markdown files for the convert-to-PDF tests."""
    import docx, openpyxl, pptx
    from pptx.util import Inches, Pt
    d = docx.Document()
    d.add_heading("Guest welcome letter", 1)
    p = d.add_paragraph("Dear Ms. Kavya Rao, thank you for choosing ")
    p.add_run("Grand Hotel Aurora").bold = True
    p.add_run(". Breakfast is served from 7 am.")
    t = d.add_table(rows=3, cols=2)
    for r, (k, v) in enumerate([("Room", "Rate"), ("Deluxe", "8,500"), ("Suite", "14,000")]):
        t.cell(r, 0).text, t.cell(r, 1).text = k, v
    d.save(folder / "letter.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "March"
    ws.append(["Date", "Rooms sold", "Revenue"])
    ws.append(["01-03", 42, 351000])
    ws.append(["02-03", 39, 327500])
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    wb.save(folder / "occupancy.xlsx")

    deck = pptx.Presentation()
    s1 = deck.slides.add_slide(deck.slide_layouts[0])
    s1.shapes.title.text = "Quarterly review"
    s1.placeholders[1].text = "Front office team"
    s2 = deck.slides.add_slide(deck.slide_layouts[1])
    s2.shapes.title.text = "Highlights"
    body = s2.placeholders[1].text_frame
    body.text = "Occupancy up to 81 percent"
    body.add_paragraph().text = "Guest score 4.6 out of 5"
    box = s2.shapes.add_textbox(Inches(6), Inches(6), Inches(3), Inches(0.6))
    box.text_frame.text = "Prepared by Meera"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    deck.save(folder / "review.pptx")

    (folder / "rates.csv").write_text("Room,Rate,Breakfast\nDeluxe,8500,yes\nSuite,\"14,000\",yes\n")
    (folder / "notes.md").write_text("# Housekeeping notes\n\n- Change towels **daily**\n- Restock the minibar\n")


def table_pdf(path):
    """A simple invoice table, for PDF to Excel."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 12)
    rows = [("Item", "Qty", "Amount"), ("Room night", "3", "25,500"), ("Breakfast", "3", "1,800"), ("Laundry", "1", "450")]
    for i, (a, b, d) in enumerate(rows):
        y = 760 - i * 22
        c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 12)
        c.drawString(72, y, a)
        c.drawRightString(300, y, b)
        c.drawRightString(450, y, d)
    c.showPage()
    c.save()


sample(out / "sample.pdf")
office(out)
table_pdf(out / "table.pdf")
photo(out / "photo.pdf")
scan(out / "scan.pdf")
form(out / "form.pdf")
subprocess.run(["qpdf", "--encrypt", "", "owner-secret", "256", "--modify=none", "--",
                str(out / "sample.pdf"), str(out / "owner-locked.pdf")], check=True)
subprocess.run(["qpdf", "--encrypt", "open-sesame", "owner-secret", "256", "--",
                str(out / "sample.pdf"), str(out / "password.pdf")], check=True)
print("fixtures written to", out)
